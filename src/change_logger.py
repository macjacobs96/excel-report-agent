"""
ChangeLogger —— 变更日志工具

记录每次生成时的修改，支持 JSON 和 Excel 两种输出格式。
"""

import json
from datetime import datetime
from pathlib import Path
from typing import Any, Dict, List

import openpyxl
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side


class ChangeLogger:
    """变更日志记录器"""

    def export(
        self,
        changes: List[Dict[str, Any]],
        output_path: Path,
        format: str = "excel",
    ) -> None:
        """
        导出变更日志。

        Args:
            changes: 变更记录列表
            output_path: 输出文件路径
            format: 输出格式 "excel" 或 "json"
        """
        if format == "json":
            self._export_json(changes, output_path)
        else:
            self._export_excel(changes, output_path)

    def _export_json(self, changes: List[Dict[str, Any]], output_path: Path) -> None:
        """导出为 JSON 格式"""
        report = {
            "generated_at": datetime.now().isoformat(),
            "total_changes": len(changes),
            "success_count": sum(1 for c in changes if c.get("status") == "SUCCESS"),
            "skipped_count": sum(1 for c in changes if "SKIPPED" in str(c.get("status", ""))),
            "failed_count": sum(1 for c in changes if "FAIL" in str(c.get("status", ""))),
            "changes": changes,
        }

        output_path.parent.mkdir(parents=True, exist_ok=True)
        with open(output_path, "w", encoding="utf-8") as f:
            json.dump(report, f, ensure_ascii=False, indent=2)

    def _export_excel(self, changes: List[Dict[str, Any]], output_path: Path) -> None:
        """导出为 Excel 格式"""
        wb = openpyxl.Workbook()
        ws = wb.active
        ws.title = "变更日志"

        # 表头样式
        header_font = Font(bold=True, size=11)
        header_fill = PatternFill(start_color="4472C4", end_color="4472C4", fill_type="solid")
        header_font_white = Font(bold=True, size=11, color="FFFFFF")
        header_alignment = Alignment(horizontal="center", vertical="center")
        thin_border = Border(
            left=Side(style="thin"),
            right=Side(style="thin"),
            top=Side(style="thin"),
            bottom=Side(style="thin"),
        )

        # 写入标题
        title_row = ["Sheet", "Cell", "字段", "原值", "新值", "状态"]
        for col, title in enumerate(title_row, 1):
            cell = ws.cell(row=1, column=col, value=title)
            cell.font = header_font_white
            cell.fill = header_fill
            cell.alignment = header_alignment
            cell.border = thin_border

        # 状态颜色
        success_fill = PatternFill(start_color="C6EFCE", end_color="C6EFCE", fill_type="solid")
        fail_fill = PatternFill(start_color="FFC7CE", end_color="FFC7CE", fill_type="solid")
        skip_fill = PatternFill(start_color="FFEB9C", end_color="FFEB9C", fill_type="solid")

        # 写入数据
        for row_idx, change in enumerate(changes, 2):
            values = [
                change.get("sheet", ""),
                change.get("cell", ""),
                change.get("field", ""),
                change.get("old_value", ""),
                change.get("new_value", ""),
                change.get("status", ""),
            ]

            for col, val in enumerate(values, 1):
                cell = ws.cell(row=row_idx, column=col, value=val)
                cell.border = thin_border
                cell.alignment = Alignment(vertical="center")

                # 状态列着色
                if col == 6:
                    cell.alignment = Alignment(horizontal="center", vertical="center")
                    status = str(val).upper()
                    if "SUCCESS" in status:
                        cell.fill = success_fill
                    elif "FAIL" in status:
                        cell.fill = fail_fill
                    elif "SKIPPED" in status:
                        cell.fill = skip_fill

        # 列宽
        col_widths = [18, 12, 20, 30, 30, 20]
        for i, width in enumerate(col_widths, 1):
            ws.column_dimensions[openpyxl.utils.get_column_letter(i)].width = width

        # 冻结表头
        ws.freeze_panes = "A2"

        # 汇总信息
        ws_summary = wb.create_sheet("汇总")
        summary_data = [
            ("生成时间", datetime.now().strftime("%Y-%m-%d %H:%M:%S")),
            ("总修改数", len(changes)),
            ("成功", sum(1 for c in changes if c.get("status") == "SUCCESS")),
            ("跳过", sum(1 for c in changes if "SKIPPED" in str(c.get("status", "")))),
            ("失败", sum(1 for c in changes if "FAIL" in str(c.get("status", "")))),
        ]

        for row_idx, (label, value) in enumerate(summary_data, 1):
            ws_summary.cell(row=row_idx, column=1, value=label).font = Font(bold=True)
            ws_summary.cell(row=row_idx, column=2, value=value)

        output_path.parent.mkdir(parents=True, exist_ok=True)
        wb.save(output_path)
        wb.close()
